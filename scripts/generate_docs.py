import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_code_block(doc, code_text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_background(cell, "F3F4F6")
    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(code_text)
    run.font.name = "Consolas"
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor(31, 41, 55) # Dark gray
    
    # Add small spacer paragraph after table
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(6)

def add_screenshot_box(doc, step_num, title, target, capture_desc):
    table = doc.add_table(rows=2, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    # Header row
    cell_hdr = table.cell(0, 0)
    set_cell_background(cell_hdr, "1E3A8A") # Navy Blue
    set_cell_margins(cell_hdr, top=100, bottom=100, left=150, right=150)
    p_hdr = cell_hdr.paragraphs[0]
    p_hdr.paragraph_format.space_before = Pt(2)
    p_hdr.paragraph_format.space_after = Pt(2)
    run_hdr = p_hdr.add_run(f"📸 SCREENSHOT #{step_num}: {title.upper()}")
    run_hdr.font.name = "Arial"
    run_hdr.font.size = Pt(10)
    run_hdr.font.bold = True
    run_hdr.font.color.rgb = RGBColor(255, 255, 255)
    
    # Body row (Placeholder area)
    cell_body = table.cell(1, 0)
    set_cell_background(cell_body, "F9FAFB") # Very light gray
    set_cell_margins(cell_body, top=180, bottom=180, left=150, right=150)
    
    p1 = cell_body.paragraphs[0]
    p1.paragraph_format.space_before = Pt(2)
    p1.paragraph_format.space_after = Pt(3)
    r1 = p1.add_run(f"Target / URL: ")
    r1.font.bold = True
    r1.font.size = Pt(9.5)
    p1.add_run(target).font.size = Pt(9.5)
    
    p2 = cell_body.add_paragraph()
    p2.paragraph_format.space_before = Pt(2)
    p2.paragraph_format.space_after = Pt(14)
    r2 = p2.add_run(f"What to capture: ")
    r2.font.bold = True
    r2.font.size = Pt(9.5)
    p2.add_run(capture_desc).font.size = Pt(9.5)
    
    p_box = cell_body.add_paragraph()
    p_box.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_box.paragraph_format.space_before = Pt(20)
    p_box.paragraph_format.space_after = Pt(20)
    r_box = p_box.add_run(f"[ PASTE SCREENSHOT #{step_num} HERE ]")
    r_box.font.name = "Arial"
    r_box.font.size = Pt(11)
    r_box.font.bold = True
    r_box.font.color.rgb = RGBColor(156, 163, 175) # Muted Gray
    
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(8)

def format_table(table, col_widths, col_alignments=None):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(table.rows):
        # Prevent row split across pages
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        
        # Header repeat
        if i == 0:
            trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
            
        for j, cell in enumerate(row.cells):
            cell.width = col_widths[j]
            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
            if i == 0:
                set_cell_background(cell, "1E3A8A")
                for p in cell.paragraphs:
                    p.paragraph_format.space_before = Pt(2)
                    p.paragraph_format.space_after = Pt(2)
                    for run in p.runs:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
                        run.font.size = Pt(9.5)
            else:
                if i % 2 == 1:
                    set_cell_background(cell, "FFFFFF")
                else:
                    set_cell_background(cell, "F9FAFB")
                for p in cell.paragraphs:
                    p.paragraph_format.space_before = Pt(2)
                    p.paragraph_format.space_after = Pt(2)
                    for run in p.runs:
                        run.font.size = Pt(9)

def main():
    doc = Document()
    
    # Page setup (Letter, 0.8 inch margins)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        
    # Title
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(4)
    run_title = p_title.add_run("EastWest Bank (EWB) Standing Order Platform")
    run_title.font.name = "Arial"
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(30, 58, 138) # Navy
    
    # Subtitle
    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(14)
    run_sub = p_sub.add_run("System Architecture, Security Perimeter & Verification Documentation")
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(13)
    run_sub.font.color.rgb = RGBColor(75, 85, 99)
    
    # Metadata callout
    p_meta = doc.add_paragraph()
    p_meta.paragraph_format.space_before = Pt(0)
    p_meta.paragraph_format.space_after = Pt(18)
    r_meta = p_meta.add_run("Member 4 Deliverable  |  Role: Platform Architect, Gateway/Security & Acceptance Lead  |  LTS 2026")
    r_meta.font.name = "Arial"
    r_meta.font.size = Pt(9.5)
    r_meta.font.italic = True
    r_meta.font.color.rgb = RGBColor(107, 114, 128)
    
    # Divider line
    doc.add_paragraph().paragraph_format.space_after = Pt(10)
    
    # --------------------------------------------------------------------------
    # 1. Executive Summary & Microservices Topology
    # --------------------------------------------------------------------------
    h1 = doc.add_heading("1. Executive Summary & Microservices Topology", level=1)
    h1.paragraph_format.space_before = Pt(12)
    h1.paragraph_format.space_after = Pt(6)
    for run in h1.runs:
        run.font.name = "Arial"
        run.font.color.rgb = RGBColor(30, 58, 138)
        
    p_desc = doc.add_paragraph(
        "The EWB Standing Order Platform is a resilient microservices system developed using Java 21 LTS, "
        "Spring Boot 3.3.4, and Spring Cloud 2023.0.3. It coordinates automated monthly transfers, enforces "
        "strict idempotency against Core Banking systems, prevents duplicate debits, and provides asynchronous, "
        "deduplicated notification delivery."
    )
    p_desc.paragraph_format.line_spacing = 1.15
    p_desc.paragraph_format.space_after = Pt(8)
    
    p_lead = doc.add_paragraph(
        "Member 4 is responsible for the foundational infrastructure, perimeter security, reverse proxy routing, "
        "asynchronous notification consumer, multi-container Docker orchestration, and integration testing."
    )
    p_lead.paragraph_format.line_spacing = 1.15
    p_lead.paragraph_format.space_after = Pt(10)

    # Topology Table
    tbl = doc.add_table(rows=8, cols=5)
    headers = ["Service Name", "Port", "Database", "Core Responsibilities", "Owner"]
    for col_idx, text in enumerate(headers):
        tbl.cell(0, col_idx).paragraphs[0].text = text
        
    data = [
        ("config-server", "8888", "Native Git/YAML", "Centralized configuration management for all services", "Member 4"),
        ("eureka-server", "8761", "In-Memory Registry", "Netflix Eureka dynamic service discovery & heartbeats", "Member 4"),
        ("gateway-service", "8080", "Stateless / JWT", "Mock IdP (/auth/login), RBAC, reverse proxy, /internal block", "Member 4"),
        ("standing-order-service", "8081", "standing_order_db", "Standing order lifecycle, versioning, schedule calendar", "Member 1"),
        ("execution-service", "8082", "execution_db", "Due order discovery, worker lease claiming, outbox relay", "Member 3"),
        ("payment-service", "8083", "payment_db", "Mock Core Banking: atomic double-entry ledger, idempotency", "Member 2"),
        ("notification-service", "8084", "notification_db", "Outbox consumer, event deduplication, outage simulation", "Member 4"),
    ]
    for row_idx, row_data in enumerate(data, start=1):
        for col_idx, val in enumerate(row_data):
            tbl.cell(row_idx, col_idx).paragraphs[0].text = val
            
    format_table(tbl, [Inches(1.5), Inches(0.6), Inches(1.3), Inches(2.6), Inches(0.9)])
    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # --------------------------------------------------------------------------
    # 2. System Personas & Authentication (contracts.md Section 2)
    # --------------------------------------------------------------------------
    h2_auth = doc.add_heading("2. Authentication & System Personas (contracts.md Section 2)", level=1)
    h2_auth.paragraph_format.space_before = Pt(12)
    h2_auth.paragraph_format.space_after = Pt(6)
    for run in h2_auth.runs:
        run.font.name = "Arial"
        run.font.color.rgb = RGBColor(30, 58, 138)
        
    p_auth_desc = doc.add_paragraph(
        "The system incorporates a Mock Identity Provider (IdP) within gateway-service on port 8080. "
        "Clients authenticate via POST /auth/login, receiving a cryptographically signed HMAC-SHA256 JWT containing "
        "their user ID, role, and authorized account list."
    )
    p_auth_desc.paragraph_format.line_spacing = 1.15
    p_auth_desc.paragraph_format.space_after = Pt(10)

    # Personas Table
    tbl_p = doc.add_table(rows=7, cols=5)
    headers_p = ["Username", "Full Name", "System Role", "Primary Account ID", "Initial Balance"]
    for col_idx, text in enumerate(headers_p):
        tbl_p.cell(0, col_idx).paragraphs[0].text = text
        
    personas_data = [
        ("asuna", "Asuna Yuuki", "ROLE_CUSTOMER", "EWB-ASU-1001", "₱20,000.00 (Active)"),
        ("kirito", "Kazuto Kirigaya", "ROLE_CUSTOMER", "EWB-KIR-5001", "₱50,000.00 (Active)"),
        ("klein", "Ryoutarou Tsuboi", "ROLE_CUSTOMER", "EWB-KLN-4001", "₱2,000.00 (Low Balance)"),
        ("heathcliff", "Akihiko Kayaba", "ROLE_CUSTOMER", "EWB-HTH-3001", "₱10,000.00 (FROZEN)"),
        ("agil", "Andrew Gilbert Mills", "ROLE_OPERATIONS", "N/A", "Operations Officer"),
        ("sinon", "Shino Asada", "ROLE_AUDITOR", "N/A", "Read-Only Auditor"),
    ]
    for row_idx, row_data in enumerate(personas_data, start=1):
        for col_idx, val in enumerate(row_data):
            tbl_p.cell(row_idx, col_idx).paragraphs[0].text = val
            
    format_table(tbl_p, [Inches(1.0), Inches(1.5), Inches(1.5), Inches(1.4), Inches(1.5)])
    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Downstream Forwarded Headers
    p_fwd_title = doc.add_paragraph()
    r_fwd = p_fwd_title.add_run("Downstream Forwarded HTTP Headers:")
    r_fwd.font.bold = True
    p_fwd_desc = doc.add_paragraph(
        "Upon validating the JWT at the API Gateway perimeter, ReverseProxyFilter extracts the identity claims "
        "and injects the following headers before dispatching the request to downstream internal microservices:\n"
        "• X-User-Id: Authenticated username (e.g. asuna)\n"
        "• X-User-Role: Assigned authorization role (e.g. ROLE_CUSTOMER)\n"
        "• X-User-Accounts: Comma-separated list of customer-owned accounts (e.g. EWB-ASU-1001,EWB-ASU-2001)"
    )
    p_fwd_desc.paragraph_format.line_spacing = 1.15
    p_fwd_desc.paragraph_format.space_after = Pt(14)

    # --------------------------------------------------------------------------
    # 3. Step-by-Step Manual Verification Guide & Screenshots
    # --------------------------------------------------------------------------
    h3_steps = doc.add_heading("3. Step-by-Step Verification Guide & Screenshot Submission", level=1)
    h3_steps.paragraph_format.space_before = Pt(14)
    h3_steps.paragraph_format.space_after = Pt(6)
    for run in h3_steps.runs:
        run.font.name = "Arial"
        run.font.color.rgb = RGBColor(30, 58, 138)
        
    p_steps_intro = doc.add_paragraph(
        "Follow these exact steps to verify the entire system end-to-end. "
        "Each section includes the command/URL, expected output, and a formatted placeholder box to insert your screenshot."
    )
    p_steps_intro.paragraph_format.line_spacing = 1.15
    p_steps_intro.paragraph_format.space_after = Pt(10)

    # Step 1
    doc.add_heading("Step 1: Automated Unit & Integration Test Suite", level=2)
    doc.add_paragraph(
        "Execute Maven across the entire multi-module project to verify that all common contracts, gateway security "
        "rules, mock IdP, and notification deduplication tests pass with 100% success."
    )
    add_code_block(doc, "mvn test")
    doc.add_paragraph("Expected Output: Reactor Summary showing BUILD SUCCESS across all 9 modules (15/15 tests passing).")
    add_screenshot_box(
        doc,
        step_num=1,
        title="Maven Automated Test Suite Success (BUILD SUCCESS)",
        target="Terminal / PowerShell",
        capture_desc="Terminal window displaying 'Reactor Summary' with all 9 modules SUCCESS and 0 test failures."
    )

    # Step 2
    doc.add_heading("Step 2: Spring Cloud Eureka Service Discovery Registry", level=2)
    doc.add_paragraph(
        "Start Eureka discovery server on port 8761 and inspect the web console in your browser."
    )
    add_code_block(doc, "mvn spring-boot:run -pl eureka-server\n# Open Browser: http://localhost:8761")
    doc.add_paragraph("Expected Output: Spring Cloud Eureka dashboard showing System Status and Registered Instances.")
    add_screenshot_box(
        doc,
        step_num=2,
        title="Eureka Discovery Dashboard at http://localhost:8761",
        target="Browser: http://localhost:8761",
        capture_desc="Eureka web dashboard displaying active server status and registered microservices."
    )

    # Step 3
    doc.add_heading("Step 3: Spring Cloud Centralized Config Server", level=2)
    doc.add_paragraph(
        "Verify that config-server on port 8888 is serving centralized YAML configuration profiles natively from config-repo/."
    )
    add_code_block(doc, "Invoke-RestMethod -Uri 'http://localhost:8888/gateway-service/default' | ConvertTo-Json -Depth 4")
    doc.add_paragraph("Expected Output: JSON payload containing propertySources with ewb.routes mappings.")
    add_screenshot_box(
        doc,
        step_num=3,
        title="Config Server Native Profile Distribution",
        target="Browser or Terminal: http://localhost:8888/gateway-service/default",
        capture_desc="JSON response confirming config-server is serving gateway-service.yml configuration."
    )

    # Step 4
    doc.add_heading("Step 4: Mock Identity Provider Login & JWT Generation", level=2)
    doc.add_paragraph(
        "Submit a POST request to /auth/login with username 'asuna' to verify JWT generation matching contracts.md Section 2."
    )
    add_code_block(doc, 
        "$body = @{ username = 'asuna' } | ConvertTo-Json\n"
        "Invoke-RestMethod -Uri 'http://localhost:8080/auth/login' -Method Post -ContentType 'application/json' -Body $body"
    )
    doc.add_paragraph("Expected Output: HTTP 200 OK with token, username 'asuna', role 'ROLE_CUSTOMER', and accounts ['EWB-ASU-1001', 'EWB-ASU-2001'].")
    add_screenshot_box(
        doc,
        step_num=4,
        title="Mock IdP Login for Asuna (ROLE_CUSTOMER)",
        target="Postman or Terminal: POST http://localhost:8080/auth/login",
        capture_desc="HTTP 200 response displaying signed JWT token and account assignments for Asuna Yuuki."
    )

    # Step 5
    doc.add_heading("Step 5: Unknown User Login Rejection (401 Unauthorized)", level=2)
    doc.add_paragraph(
        "Attempt to authenticate with an unlisted persona to ensure unauthorized credentials are rejected."
    )
    add_code_block(doc, 
        "$body = @{ username = 'unknown_hacker' } | ConvertTo-Json\n"
        "Invoke-RestMethod -Uri 'http://localhost:8080/auth/login' -Method Post -ContentType 'application/json' -Body $body"
    )
    doc.add_paragraph("Expected Output: HTTP 401 Unauthorized with status UNAUTHORIZED and message 'Unknown persona: unknown_hacker'.")
    add_screenshot_box(
        doc,
        step_num=5,
        title="Mock IdP Rejection of Unknown Persona (401 Unauthorized)",
        target="Postman or Terminal: POST http://localhost:8080/auth/login",
        capture_desc="HTTP 401 Unauthorized error response confirming perimeter authentication enforcement."
    )

    # Step 6
    doc.add_heading("Step 6: Gateway Security - Direct Internal Endpoint Block", level=2)
    doc.add_paragraph(
        "Attempt to call an internal inter-service endpoint (/internal/**) through the external API Gateway."
    )
    add_code_block(doc, "Invoke-RestMethod -Uri 'http://localhost:8080/internal/standing-orders/due' -Method Get")
    doc.add_paragraph("Expected Output: HTTP 403 Forbidden with message 'Direct access to internal endpoints is forbidden via Gateway'.")
    add_screenshot_box(
        doc,
        step_num=6,
        title="Perimeter Block of External /internal/** Access (403 Forbidden)",
        target="Postman or Terminal: GET http://localhost:8080/internal/standing-orders/due",
        capture_desc="HTTP 403 Forbidden response proving the Gateway strictly protects internal service routes."
    )

    # Step 7
    doc.add_heading("Step 7: Gateway Security - Account Ownership Validation", level=2)
    doc.add_paragraph(
        "Asuna attempts to create a standing order specifying Kirito's account (EWB-KIR-5001) as the source account."
    )
    add_code_block(doc, 
        "$hijackOrder = @{\n"
        "    sourceAccountId = 'EWB-KIR-5001'\n"
        "    destinationAccountId = 'EWB-ASU-2001'\n"
        "    amount = 5000.00\n"
        "} | ConvertTo-Json\n"
        "Invoke-RestMethod -Uri 'http://localhost:8080/standing-orders' -Method Post -ContentType 'application/json' `\n"
        "    -Headers @{ Authorization = \"Bearer $asunaToken\" } -Body $hijackOrder"
    )
    doc.add_paragraph("Expected Output: HTTP 403 Forbidden stating source account does not belong to the authenticated customer.")
    add_screenshot_box(
        doc,
        step_num=7,
        title="Source Account Hijacking Rejection (403 Forbidden)",
        target="Postman or Terminal: POST http://localhost:8080/standing-orders",
        capture_desc="HTTP 403 Forbidden response proving the Gateway checks sourceAccountId against JWT accounts."
    )

    # Step 8
    doc.add_heading("Step 8: Gateway Security - Auditor Read-Only Enforcement", level=2)
    doc.add_paragraph(
        "Sinon (ROLE_AUDITOR) attempts to perform a state-modifying POST request."
    )
    add_code_block(doc, 
        "Invoke-RestMethod -Uri 'http://localhost:8080/standing-orders' -Method Post -ContentType 'application/json' `\n"
        "    -Headers @{ Authorization = \"Bearer $sinonToken\" } -Body $hijackOrder"
    )
    doc.add_paragraph("Expected Output: HTTP 403 Forbidden stating 'Auditor role has read-only access'.")
    add_screenshot_box(
        doc,
        step_num=8,
        title="Auditor Read-Only Restriction Enforced (403 Forbidden)",
        target="Postman or Terminal: POST http://localhost:8080/standing-orders",
        capture_desc="HTTP 403 Forbidden response showing Auditor cannot perform mutating HTTP operations."
    )

    # Step 9
    doc.add_heading("Step 9: Notification Consumer - Outbox Event Ingestion", level=2)
    doc.add_paragraph(
        "Submit a new execution outbox event to POST /notifications/events matching contracts.md Section 6."
    )
    add_code_block(doc, 
        "$event = @{\n"
        "    eventId = 'evt-doc-1001'\n"
        "    executionId = 'exec-101'\n"
        "    standingOrderId = 'so-9001'\n"
        "    customerId = 'asuna'\n"
        "    eventType = 'EXECUTION_COMPLETED'\n"
        "    status = 'SUCCESS'\n"
        "    amount = 5000.00\n"
        "    currency = 'PHP'\n"
        "    sourceAccountId = 'EWB-ASU-1001'\n"
        "    destinationAccountId = 'EWB-ASU-2001'\n"
        "    paymentReference = 'TRF-20261025-001'\n"
        "    timestamp = '2026-10-25T01:00:03Z'\n"
        "} | ConvertTo-Json\n"
        "Invoke-WebRequest -Uri 'http://localhost:8084/notifications/events' -Method Post -ContentType 'application/json' -Body $event -UseBasicParsing"
    )
    doc.add_paragraph("Expected Output: HTTP 202 Accepted with status DELIVERED and deliveryId (e.g. notif-501).")
    add_screenshot_box(
        doc,
        step_num=9,
        title="Notification Event Ingestion Success (202 Accepted)",
        target="Postman or Terminal: POST http://localhost:8084/notifications/events",
        capture_desc="HTTP 202 Accepted response showing successful notification processing and delivery tracking."
    )

    # Step 10
    doc.add_heading("Step 10: Notification Deduplication - Duplicate Event Handling", level=2)
    doc.add_paragraph(
        "Re-submit the exact same event payload with eventId = 'evt-doc-1001' to verify idempotent deduplication."
    )
    add_code_block(doc, 
        "Invoke-WebRequest -Uri 'http://localhost:8084/notifications/events' -Method Post -ContentType 'application/json' -Body $event -UseBasicParsing"
    )
    doc.add_paragraph("Expected Output: HTTP 200 OK with status SKIPPED and message 'Duplicate event ID'.")
    add_screenshot_box(
        doc,
        step_num=10,
        title="Notification Deduplication Skipping Duplicate Event (200 OK SKIPPED)",
        target="Postman or Terminal: POST http://localhost:8084/notifications/events",
        capture_desc="HTTP 200 OK response with status SKIPPED confirming duplicate events are not re-delivered."
    )

    # Step 11
    doc.add_heading("Step 11: Notification Outage Simulation", level=2)
    doc.add_paragraph(
        "Submit an event with the X-Simulate-Outage: true header to simulate downstream notification gateway failure."
    )
    add_code_block(doc, 
        "Invoke-WebRequest -Uri 'http://localhost:8084/notifications/events' -Method Post -ContentType 'application/json' `\n"
        "    -Headers @{ 'X-Simulate-Outage' = 'true' } -Body $event -UseBasicParsing"
    )
    doc.add_paragraph("Expected Output: HTTP 500 Internal Server Error confirming notification failure does not corrupt the payment.")
    add_screenshot_box(
        doc,
        step_num=11,
        title="Notification Outage Simulation Clean Failure (500 Error)",
        target="Postman or Terminal: POST http://localhost:8084/notifications/events",
        capture_desc="HTTP 500 response confirming downstream notification failure simulation triggers cleanly."
    )

    # Step 12
    doc.add_heading("Step 12: Automated PowerShell Live Demo Runner", level=2)
    doc.add_paragraph(
        "Execute the automated live demonstration runner script in PowerShell."
    )
    add_code_block(doc, ".\\scripts\\run-demo.ps1")
    doc.add_paragraph("Expected Output: Formatted color terminal output displaying all test scenarios passing with [PASS] tags.")
    add_screenshot_box(
        doc,
        step_num=12,
        title="Automated Demonstration Runner (run-demo.ps1)",
        target="PowerShell Terminal: .\\scripts\\run-demo.ps1",
        capture_desc="Green terminal output displaying all contracts and resilience test scenarios passing."
    )

    # Step 13
    doc.add_heading("Step 13: Postman Collection Test Suite Execution", level=2)
    doc.add_paragraph(
        "Import postman/ewb-standing-order.postman_collection.json into Postman and execute the Collection Runner."
    )
    doc.add_paragraph("Expected Output: All 15 requests in the Postman collection executing and passing with 100% green tests.")
    add_screenshot_box(
        doc,
        step_num=13,
        title="Postman Collection Runner Execution (100% Passed)",
        target="Postman App: Collection Runner",
        capture_desc="Postman runner view displaying all requests across Auth, Orders, Transfers, and Notifications passing."
    )

    # Save Word document
    output_docx = r"c:\Users\MSB83776\Documents\antigravity\day-29-neo\docs\EWB_Standing_Order_Platform_Documentation.docx"
    doc.save(output_docx)
    print(f"Successfully generated Word documentation file at: {output_docx}")

if __name__ == "__main__":
    main()
